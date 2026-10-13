from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import load_workbook
import tifffile

from utils.registry import ACQUISITIONS
from utils.storage import (
    create_corrected_store,
    create_cropped_store,
    read_corrected_well,
    read_crop,
    read_patch_level_crop,
    write_crop,
    write_patch_level_crop,
    write_float_ome_tiff,
)


def test_zarr_round_trip_and_float_ome_tiff(tmp_path: Path):
    corrected_path = tmp_path / "corrected.ome.zarr"
    corrected_group, images = create_corrected_store(
        corrected_path,
        wells=[(2, 2), (2, 3)],
        field_shape=(8, 10),
        acquisition_metadata={"acquisition_id": "test"},
        chunk_shape=(4, 5),
        compression_level=1,
    )
    expected = np.arange(36 * 8 * 10, dtype=np.float32).reshape(36, 8, 10)
    images[1] = expected
    assert np.array_equal(read_corrected_well(corrected_path, 2, 3), expected)
    assert corrected_group.attrs["multiscales"][0]["version"] == "0.4"
    assert corrected_group.attrs["multiscales"][0]["datasets"] == [{"path": "0"}]

    crop_path = tmp_path / "crop.ome.zarr"
    group = create_cropped_store(crop_path, {"plate_name": "test"})
    crop = np.arange(64, dtype=np.float32).reshape(8, 8)
    write_crop(
        group,
        "measurement_01_dapi",
        crop,
        metadata={"shift_y": 0, "shift_x": 0},
        chunk_shape=(4, 4),
        compression_level=1,
    )
    assert np.array_equal(read_crop(crop_path, "measurement_01_dapi"), crop)

    write_patch_level_crop(
        group,
        "m09_12",
        5,
        "measurement_09_dapi",
        crop,
        metadata={"reference_measurement": 9},
        chunk_shape=(4, 4),
        compression_level=1,
    )
    assert np.array_equal(
        read_patch_level_crop(
            crop_path, "m09_12", 5, "measurement_09_dapi"
        ),
        crop,
    )

    tiff_path = tmp_path / "patch.ome.tif"
    write_float_ome_tiff(
        tiff_path,
        crop,
        metadata={"acquisition_id": "measurement_01_dapi"},
        compression="zlib",
    )
    reloaded = tifffile.imread(tiff_path)
    assert reloaded.dtype == np.float32
    assert np.array_equal(reloaded, crop)


def test_clipping_template_has_exact_registry_ids():
    path = Path(__file__).resolve().parents[1] / "config" / "clip_limits.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["clip_limits"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "acquisition_id",
        "measurement",
        "source_channel",
        "marker",
        "min_16bit",
        "max_16bit",
    )
    assert [row[0] for row in rows[1:]] == [item.acquisition_id for item in ACQUISITIONS]
    assert all(row[4] is None and row[5] is None for row in rows[1:])
