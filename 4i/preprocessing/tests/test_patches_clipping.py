from __future__ import annotations

import numpy as np
from openpyxl import Workbook

from utils.clipping import clip_to_uint8, load_clip_limits
from utils.patches import patch_plan, quarantined_patch_ids
from utils.registry import ACQUISITIONS


def test_failed_field_quarantines_only_intersecting_patches():
    plan = patch_plan((0, 6144, 0, 6144), 1024)
    quarantined = quarantined_patch_ids(plan, {2})
    assert quarantined == {"patch_0001", "patch_0002", "patch_0007", "patch_0008"}
    assert len(plan) == 36
    assert len(plan) - len(quarantined) == 32


def test_uint8_clipping_endpoints_and_nonfinite_values():
    image = np.array([[-10, 0, 50, 100, 110, np.nan]], dtype=np.float32)
    result = clip_to_uint8(image, 0, 100)
    assert result.tolist() == [[0, 0, 128, 255, 255, 0]]


def test_excel_limits_validate_exact_acquisition_keys(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "clip_limits"
    sheet.append(
        [
            "acquisition_id",
            "measurement",
            "source_channel",
            "marker",
            "min_16bit",
            "max_16bit",
        ]
    )
    for acquisition in ACQUISITIONS:
        sheet.append(
            [
                acquisition.acquisition_id,
                acquisition.measurement,
                acquisition.source_channel,
                acquisition.marker,
                100,
                5000,
            ]
        )
    path = tmp_path / "limits.xlsx"
    workbook.save(path)
    limits = load_clip_limits(path)
    assert set(limits) == {item.acquisition_id for item in ACQUISITIONS}
    assert all(row["min_16bit"] == 100 and row["max_16bit"] == 5000 for row in limits.values())
    assert all(row["custom_clipping"] is True for row in limits.values())


def test_blank_excel_limits_use_full_16bit_default(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "clip_limits"
    sheet.append(
        [
            "acquisition_id",
            "measurement",
            "source_channel",
            "marker",
            "min_16bit",
            "max_16bit",
        ]
    )
    for acquisition in ACQUISITIONS:
        sheet.append(
            [
                acquisition.acquisition_id,
                acquisition.measurement,
                acquisition.source_channel,
                acquisition.marker,
                None,
                None,
            ]
        )
    path = tmp_path / "blank_limits.xlsx"
    workbook.save(path)
    limits = load_clip_limits(path)
    assert set(limits) == {item.acquisition_id for item in ACQUISITIONS}
    assert all(
        row["min_16bit"] == 0
        and row["max_16bit"] == 65535
        and row["limits_source"] == "default_full_16bit"
        and row["custom_clipping"] is False
        for row in limits.values()
    )


def test_only_one_blank_limit_is_rejected(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "clip_limits"
    sheet.append(
        [
            "acquisition_id",
            "measurement",
            "source_channel",
            "marker",
            "min_16bit",
            "max_16bit",
        ]
    )
    sheet.append(["measurement_01_dapi", 1, 1, "DAPI", None, 5000])
    path = tmp_path / "partial_limits.xlsx"
    workbook.save(path)
    import pytest

    with pytest.raises(ValueError, match="only one limit"):
        load_clip_limits(path)


def test_retained_patch_ids_are_reused_for_all_acquisitions():
    plan = patch_plan((0, 6144, 0, 6144), 1024)
    excluded = quarantined_patch_ids(plan, {2, 13})
    retained = [item["patch_id"] for item in plan if item["patch_id"] not in excluded]
    by_acquisition = {
        acquisition.acquisition_id: list(retained) for acquisition in ACQUISITIONS
    }
    assert len({tuple(ids) for ids in by_acquisition.values()}) == 1
