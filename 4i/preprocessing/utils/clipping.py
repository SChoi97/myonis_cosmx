"""Validated clipping workbook loading and deterministic uint8 conversion."""

from __future__ import annotations

from pathlib import Path

import numpy as np


REQUIRED_COLUMNS = (
    "acquisition_id",
    "measurement",
    "source_channel",
    "marker",
    "min_16bit",
    "max_16bit",
)

DEFAULT_MIN_16BIT = 0.0
DEFAULT_MAX_16BIT = 65535.0


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def load_clip_limits(path: str | Path) -> dict[str, dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ImportError("openpyxl is required to read clipping limits.") from error
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["clip_limits"] if "clip_limits" in workbook.sheetnames else workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
    except StopIteration:
        raise ValueError("Clipping workbook is empty.")
    missing = set(REQUIRED_COLUMNS) - set(headers)
    if missing:
        raise ValueError(f"Clipping workbook lacks columns {sorted(missing)}.")
    rows = {}
    for excel_row, values in enumerate(iterator, start=2):
        row = dict(zip(headers, values))
        acquisition_id = str(row.get("acquisition_id") or "").strip()
        if not acquisition_id:
            continue
        if acquisition_id in rows:
            raise ValueError(f"Duplicate acquisition_id {acquisition_id!r} in row {excel_row}.")
        minimum_value = row["min_16bit"]
        maximum_value = row["max_16bit"]
        if _blank(minimum_value) and _blank(maximum_value):
            minimum = DEFAULT_MIN_16BIT
            maximum = DEFAULT_MAX_16BIT
            limits_source = "default_full_16bit"
            custom_clipping = False
        elif _blank(minimum_value) or _blank(maximum_value):
            raise ValueError(
                f"Clipping row {excel_row} has only one limit. Leave both blank "
                "for the default 0-65535 range, or fill both values."
            )
        else:
            try:
                minimum = float(minimum_value)
                maximum = float(maximum_value)
            except (TypeError, ValueError) as error:
                raise ValueError(
                    f"Invalid min/max in clipping row {excel_row}."
                ) from error
            limits_source = "workbook"
            custom_clipping = True
        if not np.isfinite(minimum) or not np.isfinite(maximum) or maximum <= minimum:
            raise ValueError(
                f"Row {excel_row} requires finite max_16bit > min_16bit."
            )
        rows[acquisition_id] = {
            **row,
            "acquisition_id": acquisition_id,
            "min_16bit": minimum,
            "max_16bit": maximum,
            "limits_source": limits_source,
            "custom_clipping": custom_clipping,
        }
    return rows


def clip_to_uint8(image: np.ndarray, minimum: float, maximum: float) -> np.ndarray:
    if not np.isfinite(minimum) or not np.isfinite(maximum) or maximum <= minimum:
        raise ValueError("maximum must be finite and greater than minimum.")
    values = np.asarray(image, dtype=np.float32)
    scaled = np.clip((values - minimum) / (maximum - minimum), 0, 1)
    scaled[~np.isfinite(scaled)] = 0
    return np.rint(scaled * 255).astype(np.uint8)
