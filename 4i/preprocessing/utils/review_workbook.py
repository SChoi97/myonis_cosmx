"""Read the human-reviewed illumination-method workbook."""

from __future__ import annotations

from pathlib import Path


REVIEW_SHEETS = ("Marker Selections", "DAPI Selections")


def normalize_method(value: object) -> str:
    """Return the exact production method name from a friendly workbook value."""
    text = str(value or "").strip().lower().replace("_", "-")
    if text in {"basic", "basicpy", "basic-py"}:
        return "BaSiC"
    if text in {"harmony", "harmony-like", "harmonylike"}:
        return "Harmony-like"
    if not text:
        return ""
    raise ValueError(
        f"Unsupported illumination method {value!r}; use BaSiC or Harmony-like."
    )


def read_review_workbook(path: str | Path) -> dict[str, dict[str, str]]:
    """Load choices from the two flat review sheets, deriving approval from choice."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ImportError("openpyxl is required to read the review workbook.") from error

    workbook_path = Path(path).expanduser().resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    choices: dict[str, dict[str, str]] = {}
    for sheet_name in REVIEW_SHEETS:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Review workbook is missing required sheet {sheet_name!r}."
            )
        sheet = workbook[sheet_name]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        required = {"acquisition_id", "chosen_method"}
        if not required.issubset(headers):
            raise ValueError(
                f"{sheet_name!r} must contain columns {sorted(required)}."
            )
        column = {name: index for index, name in enumerate(headers)}
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            acquisition_id = str(values[column["acquisition_id"]] or "").strip()
            if not acquisition_id:
                continue
            if acquisition_id in choices:
                raise ValueError(
                    f"Duplicate acquisition_id {acquisition_id!r} in review workbook."
                )
            method = normalize_method(values[column["chosen_method"]])
            note = ""
            if "review_notes" in column:
                note = str(values[column["review_notes"]] or "").strip()
            choices[acquisition_id] = {
                "chosen_method": method,
                "review_status": "approved" if method else "pending",
                "review_notes": note,
                "review_workbook_path": str(workbook_path),
                "review_sheet": sheet_name,
                "review_row": str(row_number),
            }
    return choices
